"""Метод 1: непрерывные функции принадлежности расстояния пешком."""

from pathlib import Path
import re

import numpy as np
from openpyxl import load_workbook


BOOK = Path(__file__).with_name("Лабораторная_1_расстояние_пешком_исправленная.xlsx")
OUTPUT = Path(__file__).with_name("grafik_expertov_rasstoyanie_skfuzzy.png")
TERMS = ("Близко", "Средне", "Далеко")


def read_boundaries(labels):
    """Преобразуем подписи 0–1, 1–2, ... в границы 0, 1, 2, ... км."""
    pairs = []
    for label in labels:
        match = re.fullmatch(r"\s*(\d+(?:[.,]\d+)?)\s*[–-]\s*(\d+(?:[.,]\d+)?)\s*(?:км)?\s*", label)
        if match is None:
            raise ValueError(f"Не удалось прочитать интервал расстояния: {label!r}")
        pairs.append(tuple(float(number.replace(",", ".")) for number in match.groups()))
    if any(left >= right for left, right in pairs):
        raise ValueError("Границы каждого интервала должны возрастать.")
    if any(not np.isclose(pairs[i][1], pairs[i + 1][0]) for i in range(len(pairs) - 1)):
        raise ValueError("Интервалы расстояния должны идти подряд без пропусков.")
    return np.array([pairs[0][0], *(right for _, right in pairs)])


def trapezoid_from_votes(degrees, boundaries):
    """Находим начало, плато и конец терма по интервалам с долей голосов 0..1."""
    positive = np.flatnonzero(degrees > 0)
    full = np.flatnonzero(np.isclose(degrees, 1))
    if not len(positive) or not len(full):
        raise ValueError("Для трапециевидной функции нужен хотя бы один интервал со степенью 1.")
    if not np.array_equal(positive, np.arange(positive[0], positive[-1] + 1)):
        raise ValueError("Интервалы с ненулевой принадлежностью должны идти подряд.")
    if not np.array_equal(full, np.arange(full[0], full[-1] + 1)):
        raise ValueError("Интервалы со степенью 1 должны образовывать одно плато.")

    first_positive, last_positive = positive[0], positive[-1]
    first_full, last_full = full[0], full[-1]
    a = boundaries[first_positive]
    b = boundaries[first_full]
    c = boundaries[last_full + 1]
    d = boundaries[last_positive + 1]
    # На краях области: левое и правое плечи трапеции.
    if first_full == 0:
        a = b = boundaries[0]
    if last_full == len(degrees) - 1:
        c = d = boundaries[-1]
    return np.array([a, b, c, d])


def main():
    workbook = load_workbook(BOOK, read_only=True, data_only=True)
    try:
        sheet = workbook["Метод 1"]
        intervals = [str(sheet.cell(1, col).value) for col in range(3, 9)]
        # Пять экспертов × три терма × шесть интервалов. Только исходные 0/1.
        votes = np.array(
            [
                [
                    [sheet.cell(2 + 3 * expert + term, col).value for col in range(3, 9)]
                    for term in range(3)
                ]
                for expert in range(5)
            ],
            dtype=float,
        )
    finally:
        workbook.close()

    if not np.isin(votes, [0, 1]).all():
        raise ValueError("В исходной таблице допустимы только оценки 0 и 1.")
    if not np.all(votes.sum(axis=1) == 1):
        raise ValueError("Для каждого интервала эксперт должен выбрать ровно один терм.")

    # Доля экспертов, выбравших терм: число единиц / 5.
    memberships = votes.mean(axis=0)
    if np.any((memberships[0] > 0) & (memberships[2] > 0)):
        raise ValueError("В этих данных «Близко» и «Далеко» не должны пересекаться.")

    boundaries = read_boundaries(intervals)
    params = {
        term: trapezoid_from_votes(degrees, boundaries)
        for term, degrees in zip(TERMS, memberships)
    }

    print("Степени принадлежности по оценкам экспертов:")
    for term, values in zip(TERMS, memberships):
        print(f"{term:8s}: {values}")

    try:
        import matplotlib.pyplot as plt
        import skfuzzy as fuzz
    except ModuleNotFoundError as exc:
        raise SystemExit(
            f"Для графика не хватает пакета {exc.name}. "
            "Установите matplotlib и scikit-fuzzy в используемый Python."
        ) from exc

    # Доли голосов относятся ко всему интервалу, а не к конкретной точке внутри
    # него. Поэтому они определяют области переходов; плавные прямые строим
    # между настоящими границами интервалов, не придумывая их середины.
    x = np.linspace(boundaries[0], boundaries[-1], 601)
    curves = {term: fuzz.trapmf(x, values) for term, values in params.items()}
    if not np.allclose(sum(curves.values()), 1, atol=1e-9):
        raise ValueError("Полученные термы должны покрывать всю область: сумма степеней равна 1.")

    print("Параметры trapmf [a, b, c, d], км:")
    for term, values in params.items():
        print(f"{term:8s}: {values.tolist()}")

    fig, ax = plt.subplots(figsize=(11, 5.5))
    for term, color in zip(TERMS, ("#168054", "#e78016", "#d83d48")):
        ax.plot(x, curves[term], linewidth=2.5, color=color, label=term)
    ax.set_xlim(boundaries[0], boundaries[-1])
    ax.set_xticks(boundaries)
    ax.set_yticks(np.arange(0, 1.01, 0.2))
    ax.set_xlabel("Расстояние пешком, км")
    ax.set_ylabel("Степень принадлежности")
    ax.set_ylim(0, 1.05)
    ax.set_title("Функции принадлежности расстояния пешком")
    ax.grid(alpha=0.25)
    ax.legend()
    fig.tight_layout()
    fig.savefig(OUTPUT, dpi=180)
    print(f"График сохранён: {OUTPUT}")
    plt.show()


if __name__ == "__main__":
    main()
