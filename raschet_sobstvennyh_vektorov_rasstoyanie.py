"""Метод 2: собственные значения, векторы и графики scikit-fuzzy."""

from pathlib import Path
import sys

import numpy as np
from openpyxl import load_workbook


BOOK = Path(__file__).with_name("Лабораторная_1_расстояние_пешком_исправленная.xlsx")
TERMS = (("Близко", 28), ("Средне", 40), ("Далеко", 52))
PLOT_STYLE = {
    "Близко": ("blizko", "#168054"),
    "Средне": ("sredne", "#e78016"),
    "Далеко": ("daleko", "#d83d48"),
}
N = 6


def read_matrix(sheet, first_row):
    """Вводим верхний треугольник; снизу стоят обратные оценки 1/a_ij."""
    matrix = np.eye(N)
    for i in range(N):
        for j in range(i + 1, N):
            cell = sheet.cell(first_row + i, 2 + j)
            value = cell.value
            if not isinstance(value, (int, float)) or value <= 0:
                raise ValueError(f"В ячейке {cell.coordinate} нужна положительная оценка.")
            matrix[i, j] = value
            matrix[j, i] = 1 / value
    return matrix


def format_eigenvalue(value):
    if abs(value.imag) < 1e-10:
        return f"{value.real:.6f}"
    return f"{value.real:.6f} {value.imag:+.6f}i"


def plot_membership(intervals, memberships):
    try:
        import matplotlib.pyplot as plt
        from skfuzzy import control as ctrl
    except ModuleNotFoundError as exc:
        print(
            f"\nГрафики не построены: не хватает пакета {exc.name}. "
            "Установите matplotlib и scikit-fuzzy в используемый Python."
        )
        return

    x = np.arange(N)
    distance = ctrl.Antecedent(x, "Расстояние пешком")
    for term, _ in TERMS:
        distance[term] = memberships[term]
    figures = []

    def prepare_axes(ax, title):
        ax.set_xticks(x, intervals)
        ax.set_ylim(0, 1.05)
        ax.set_yticks(np.linspace(0, 1, 6))
        ax.set_xlabel("Интервал расстояния пешком, км")
        ax.set_ylabel("Степень принадлежности")
        ax.set_title(title)
        ax.grid(alpha=0.25)

    fig, ax = plt.subplots(figsize=(11, 5.5))
    for term, _ in TERMS:
        _, color = PLOT_STYLE[term]
        ax.plot(x, distance[term].mf, marker="o", linewidth=2, color=color, label=term)
    prepare_axes(ax, "Функции принадлежности по методу Саати")
    ax.legend()
    fig.tight_layout()
    path = BOOK.with_name("grafik_saaty_rasstoyanie_skfuzzy_obshchii.png")
    fig.savefig(path, dpi=180)
    figures.append(fig)
    print(f"Общий график: {path}")

    for term, _ in TERMS:
        slug, color = PLOT_STYLE[term]
        fig, ax = plt.subplots(figsize=(10, 5))
        ax.plot(x, distance[term].mf, marker="o", linewidth=2, color=color)
        prepare_axes(ax, f"{term}: функция принадлежности по методу Саати")
        fig.tight_layout()
        path = BOOK.with_name(f"grafik_saaty_rasstoyanie_skfuzzy_{slug}.png")
        fig.savefig(path, dpi=180)
        figures.append(fig)
        print(f"График «{term}»: {path}")

    plt.show()
    for fig in figures:
        plt.close(fig)


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    workbook = load_workbook(BOOK, read_only=True, data_only=False)
    memberships = {}
    try:
        sheet = workbook["Метод 2"]
        intervals = [str(sheet.cell(27, col).value) for col in range(2, 2 + N)]

        for term, first_row in TERMS:
            matrix = read_matrix(sheet, first_row)
            eigenvalues, eigenvectors = np.linalg.eig(matrix)
            principal_index = int(np.argmax(eigenvalues.real))
            lambda_max = eigenvalues[principal_index]
            vector = eigenvectors[:, principal_index]
            if abs(lambda_max.imag) > 1e-10 or np.max(np.abs(vector.imag)) > 1e-10:
                raise ArithmeticError("Главное собственное значение или вектор не вещественные.")
            vector = vector.real
            if vector.sum() < 0:
                vector = -vector
            if np.any(vector <= 0):
                raise ArithmeticError("Главный собственный вектор должен быть положительным.")

            weight = vector / vector.sum()  # Нормировка собственного вектора: сумма W равна 1.
            mu = weight / weight.max()      # Нормировка функции принадлежности: max μ = 1.
            memberships[term] = mu
            residual = np.max(np.abs(matrix @ weight - lambda_max.real * weight))

            print(f"\n{term.upper()}")
            print("Все собственные значения:")
            for value in sorted(eigenvalues, key=lambda x: x.real, reverse=True):
                print(" ", format_eigenvalue(value))
            print(f"λmax = {lambda_max.real:.6f}; λmax − {N} = {lambda_max.real - N:.6f}")
            print("Интервал       W          μ")
            for interval, w, membership in zip(intervals, weight, mu):
                print(f"{interval:<10} {w:.6f}   {membership:.6f}")
            print(f"Проверка AW = λmax·W: максимальная ошибка {residual:.2e}")
    finally:
        workbook.close()

    plot_membership(intervals, memberships)


if __name__ == "__main__":
    main()
